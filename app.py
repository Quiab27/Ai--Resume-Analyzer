import io
import re
import time
from xml.sax.saxutils import escape

import streamlit as st
from google import genai
from google.genai import errors, types
from pydantic import BaseModel, Field
from pypdf import PdfReader

# ---------- Config ----------
MODELS = ["gemini-3.5-flash", "gemini-3.5-flash-lite", "gemini-2.5-flash"]
TIMEOUT_MS = 120_000
MAX_ATTEMPTS = 3
RETRYABLE_CODES = {429, 500, 502, 503, 504}
MAX_RESUME_CHARS = 20_000
MAX_JD_CHARS = 10_000


# ---------- Schema ----------
class SkillGroup(BaseModel):
    category: str = Field(description="e.g. Programming Languages, AI & Data, Tools")
    items: list[str]


class Project(BaseModel):
    name: str
    tech: str = Field(description="Tools/technologies used, comma separated")
    bullets: list[str]


class Experience(BaseModel):
    title: str
    organization: str
    dates: str
    bullets: list[str]


class Education(BaseModel):
    degree: str
    institution: str
    dates: str
    details: str = Field(description="CGPA, relevant modules, etc. Empty string if none")


class Resume(BaseModel):
    name: str
    contact: list[str] = Field(description="Email, phone, city, LinkedIn, GitHub, as in the original")
    headline: str = Field(description="Short role-aligned title, e.g. 'AI / Software Engineering Intern'")
    summary: str
    education: list[Education]
    skill_groups: list[SkillGroup]
    projects: list[Project]
    experience: list[Experience]
    achievements: list[str]
    languages: str = Field(description="Spoken languages, comma separated. Empty string if none")


class Analysis(BaseModel):
    match_score: int = Field(description="Overall match from 0 to 100")
    summary: str
    strengths: list[str]
    missing_skills: list[str]
    recommendations: list[str]
    ats_keywords: list[str]
    resume: Resume


# ---------- Prompt ----------
def build_prompt(resume: str, title: str, jd: str) -> str:
    return f"""You are an expert technical recruiter and resume writer. Tailor the candidate's resume
for the job below so it reads as if written specifically for this role.

TARGET JOB TITLE: {title}

JOB DESCRIPTION:
{jd[:MAX_JD_CHARS]}

ORIGINAL RESUME:
{resume[:MAX_RESUME_CHARS]}

WRITING RULES
- TRUTH: use only facts in the original resume. Never invent employers, dates, tools, metrics,
  results or achievements. If the job wants something the candidate lacks, do NOT add it; put
  it in missing_skills. Use numbers only if they appear in the original.
- headline: a short title aligned to the target role (not a made-up job the candidate holds).
- summary: 2-3 sentences. Name the target role, then the candidate's strongest relevant evidence
  (degree focus, 2-3 most relevant skills/projects). No cliches like "passionate",
  "hard-working", "team player".
- education: keep factual. Put relevant modules in details only if they relate to the job.
- skill_groups: 3-5 categories, ordered by relevance to the job; inside each, most relevant
  skills first. Only skills present in the original. Use the job's wording where the candidate
  truly has that skill (e.g. write the job's exact term for a tool they used).
- projects: choose the 4 most relevant to this job, ordered by relevance; drop or merge the
  rest. 2-3 bullets each. Every bullet: strong past-tense action verb + what was built +
  how (tools) + outcome/purpose, one line (max ~25 words). Emphasise the parts of each project
  that match the job description. No repeated opening verbs within a project.
- experience: include tutoring/work experience, reframed toward transferable skills the job
  asks for (communication, mentoring, mathematics, etc.), 2-3 bullets each.
- achievements: only real ones from the original, most relevant first, max 3.
- Whole resume must fit one A4 page: be selective and concise.
- Also return match_score, summary, strengths, missing_skills, recommendations (concrete advice
  on what the candidate could learn or add), and ats_keywords (job keywords that now appear
  in the resume).
"""


# ---------- Gemini call ----------
def extract_pdf_text(uploaded_file) -> str:
    reader = PdfReader(uploaded_file)
    return "\n".join((p.extract_text() or "") for p in reader.pages).strip()


def analyze(api_key: str, prompt: str, status=None) -> Analysis:
    client = genai.Client(
        api_key=api_key, http_options=types.HttpOptions(timeout=TIMEOUT_MS)
    )
    config = types.GenerateContentConfig(
        response_mime_type="application/json",
        response_schema=Analysis,
        temperature=0.4,
    )
    last_err: Exception | None = None
    tried: list[str] = []
    for model in MODELS:
        for attempt in range(1, MAX_ATTEMPTS + 1):
            try:
                if status:
                    status.update(label=f"Tailoring with {model} (attempt {attempt}/{MAX_ATTEMPTS})...")
                resp = client.models.generate_content(model=model, contents=prompt, config=config)
                if resp.parsed is not None:
                    return resp.parsed
                return Analysis.model_validate_json(resp.text)
            except errors.APIError as e:
                last_err = e
                code = getattr(e, "code", None)
                if code in RETRYABLE_CODES:
                    time.sleep(2 ** attempt)
                    continue
                if code == 404:
                    tried.append(f"{model}: not available")
                    break
                raise
            except Exception as e:
                last_err = e
                time.sleep(2 ** attempt)
    raise RuntimeError(f"All models/attempts failed. Skipped: {tried or 'none'}. Last error: {last_err}")


# ---------- Export helpers ----------
_CHAR_FIX = {"\u2192": "->", "\u2022": "-", "\u2019": "'", "\u2018": "'", "\u201c": '"',
             "\u201d": '"', "\u00a0": " ", "\u2011": "-", "\u200b": ""}


def _clean(s: str) -> str:
    s = s or ""
    for a, b in _CHAR_FIX.items():
        s = s.replace(a, b)
    return s.strip()


def _slug(s: str) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", s).strip("_")[:40] or "resume"


def resume_to_markdown(r: Resume) -> str:
    out = [f"# {r.name}", f"**{r.headline}**", " | ".join(r.contact), "",
           "## Summary", r.summary, "", "## Education"]
    for e in r.education:
        out.append(f"**{e.degree}**, {e.institution} ({e.dates})")
        if e.details:
            out.append(f"{e.details}")
        out.append("")
    out.append("## Skills")
    for g in r.skill_groups:
        out.append(f"- **{g.category}:** {', '.join(g.items)}")
    out += ["", "## Projects"]
    for p in r.projects:
        out.append(f"**{p.name}** | *{p.tech}*")
        out += [f"- {b}" for b in p.bullets]
        out.append("")
    if r.experience:
        out.append("## Experience")
        for x in r.experience:
            out.append(f"**{x.title}**, {x.organization} ({x.dates})")
            out += [f"- {b}" for b in x.bullets]
            out.append("")
    if r.achievements:
        out.append("## Achievements")
        out += [f"- {a}" for a in r.achievements]
        out.append("")
    if r.languages:
        out += ["## Languages", r.languages]
    return "\n".join(out)


# ----- Word (.docx) -----
def build_docx(r: Resume) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    NAVY = RGBColor(0x1F, 0x3A, 0x5F)
    GREY = RGBColor(0x55, 0x55, 0x55)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(15)
    sec.top_margin = sec.bottom_margin = Mm(12)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)

    def heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = NAVY
        pPr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "1F3A5F")):
            bottom.set(qn(k), v)
        borders.append(bottom)
        pPr.append(borders)

    def split_line(bold_text, rest, right):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(Mm(180), WD_TAB_ALIGNMENT.RIGHT)
        b = p.add_run(_clean(bold_text))
        b.bold = True
        if rest:
            r2 = p.add_run(rest[: len(rest) - len(rest.lstrip())] + _clean(rest))
            r2.italic = True
            r2.font.color.rgb = GREY
        if right:
            p.add_run("\t" + _clean(right)).font.color.rgb = GREY

    def bullet(text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Mm(5)
        p.paragraph_format.first_line_indent = Mm(-4)
        p.add_run(_clean(text))

    # Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(_clean(r.name).upper())
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(_clean(r.headline))
    run.font.size = Pt(11)
    run.bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("  |  ".join(_clean(c) for c in r.contact if c.strip()))
    run.font.size = Pt(9)
    run.font.color.rgb = GREY

    heading("Professional Summary")
    doc.add_paragraph(_clean(r.summary))

    heading("Education")
    for e in r.education:
        split_line(e.degree, f", {e.institution}", e.dates)
        if e.details.strip():
            doc.add_paragraph(_clean(e.details))

    heading("Technical Skills")
    for g in r.skill_groups:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.add_run(_clean(g.category) + ": ").bold = True
        p.add_run(_clean(", ".join(g.items)))

    heading("Projects")
    for pr in r.projects:
        split_line(pr.name, f"  |  {pr.tech}" if pr.tech.strip() else "", "")
        for b in pr.bullets:
            bullet(b)

    if r.experience:
        heading("Experience")
        for x in r.experience:
            split_line(x.title, f", {x.organization}", x.dates)
            for b in x.bullets:
                bullet(b)

    if r.achievements:
        heading("Achievements")
        for a in r.achievements:
            bullet(a)

    if r.languages.strip():
        heading("Languages")
        doc.add_paragraph(_clean(r.languages))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----- PDF -----
def build_pdf(r: Resume) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Table, TableStyle

    NAVY = colors.HexColor("#1F3A5F")
    GREY = colors.HexColor("#555555")
    base = ParagraphStyle("base", fontName="Helvetica", fontSize=9.5, leading=12.2)
    name_st = ParagraphStyle("name", parent=base, fontName="Helvetica-Bold", fontSize=20,
                             leading=24, alignment=TA_CENTER, textColor=NAVY)
    head_st = ParagraphStyle("headline", parent=base, fontName="Helvetica-Bold", fontSize=11,
                             leading=14, alignment=TA_CENTER)
    contact_st = ParagraphStyle("contact", parent=base, fontSize=8.5, alignment=TA_CENTER, textColor=GREY)
    sec_st = ParagraphStyle("sec", parent=base, fontName="Helvetica-Bold", fontSize=10.5,
                            textColor=NAVY, spaceBefore=7, spaceAfter=1)
    bullet_st = ParagraphStyle("bullet", parent=base, leftIndent=11, bulletIndent=2, spaceAfter=1)
    left_st = ParagraphStyle("left", parent=base)
    right_st = ParagraphStyle("right", parent=base, alignment=2, textColor=GREY)

    def esc(s):
        return escape(_clean(s))

    story = []

    def section(title):
        story.append(Paragraph(title.upper(), sec_st))
        story.append(HRFlowable(width="100%", thickness=0.8, color=NAVY, spaceAfter=3))

    def split_row(bold_text, rest, right):
        left = f"<b>{esc(bold_text)}</b>"
        if rest:
            left += f'<font color="#555555"><i>{" " * (len(rest) - len(rest.lstrip()))}{esc(rest)}</i></font>'
        t = Table([[Paragraph(left, left_st), Paragraph(esc(right), right_st)]],
                  colWidths=[135 * mm, 45 * mm])
        t.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                               ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
                               ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    story.append(Paragraph(esc(r.name).upper(), name_st))
    story.append(Paragraph(esc(r.headline), head_st))
    story.append(Paragraph("  |  ".join(esc(c) for c in r.contact if c.strip()), contact_st))

    section("Professional Summary")
    story.append(Paragraph(esc(r.summary), base))

    section("Education")
    for e in r.education:
        story.append(split_row(e.degree, f", {e.institution}", e.dates))
        if e.details.strip():
            story.append(Paragraph(esc(e.details), base))

    section("Technical Skills")
    for g in r.skill_groups:
        story.append(Paragraph(f"<b>{esc(g.category)}:</b> {esc(', '.join(g.items))}", ParagraphStyle(
            "sk", parent=base, spaceAfter=1)))

    section("Projects")
    for pr in r.projects:
        block = [split_row(pr.name, f"  |  {pr.tech}" if pr.tech.strip() else "", "")]
        block += [Paragraph(esc(b), bullet_st, bulletText="\u2022") for b in pr.bullets]
        story.append(KeepTogether(block))

    if r.experience:
        section("Experience")
        for x in r.experience:
            block = [split_row(x.title, f", {x.organization}", x.dates)]
            block += [Paragraph(esc(b), bullet_st, bulletText="\u2022") for b in x.bullets]
            story.append(KeepTogether(block))

    if r.achievements:
        section("Achievements")
        for a in r.achievements:
            story.append(Paragraph(esc(a), bullet_st, bulletText="\u2022"))

    if r.languages.strip():
        section("Languages")
        story.append(Paragraph(esc(r.languages), base))

    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
                      topMargin=12 * mm, bottomMargin=12 * mm, title=f"{r.name} - Resume",
                      author=r.name).build(story)
    return buf.getvalue()


# ---------- UI ----------
import html as _html
import os

st.set_page_config(page_title="AI Resume Tailor", page_icon="📄", layout="wide")

CSS = """
<style>
#MainMenu, footer, header[data-testid="stHeader"] {visibility: hidden; height: 0;}
.block-container {padding-top: 1.6rem; max-width: 1150px;}
.hero {
  background: linear-gradient(135deg, #1F3A5F 0%, #4F46E5 100%);
  border-radius: 20px; padding: 34px 38px; color: #fff; margin-bottom: 22px;
  box-shadow: 0 10px 30px rgba(79,70,229,.25);
}
.hero h1 {margin: 0 0 6px 0; font-size: 2.1rem; font-weight: 800; color: #fff; padding: 0;}
.hero p {margin: 0; font-size: 1.05rem; opacity: .9;}
.hero .pill {display:inline-block; background: rgba(255,255,255,.18); padding: 4px 12px;
  border-radius: 999px; font-size: .8rem; margin: 14px 8px 0 0;}
.step-title {font-weight: 700; font-size: 1.05rem; color: #1F3A5F; margin-bottom: 2px;}
.step-badge {display:inline-block; width: 26px; height: 26px; line-height: 26px; text-align:center;
  background:#4F46E5; color:#fff; border-radius:50%; font-size:.85rem; margin-right:8px;}
div[data-testid="stVerticalBlockBorderWrapper"] {border-radius: 16px;}
.stButton > button[kind="primary"], .stDownloadButton > button {
  border-radius: 12px; font-weight: 600; padding: .6rem 1.4rem; border: none;
}
.stButton > button[kind="primary"] {
  background: linear-gradient(135deg, #4F46E5, #1F3A5F); color: #fff; width: 100%;
  font-size: 1.05rem; padding: .75rem 1rem; box-shadow: 0 6px 16px rgba(79,70,229,.3);
}
.stButton > button[kind="primary"]:hover {filter: brightness(1.08);}
.stDownloadButton > button {background:#fff; border: 1.5px solid #4F46E5; color:#4F46E5;}
.stDownloadButton > button:hover {background:#4F46E5; color:#fff;}
.ring {--p: 0; --c: #16a34a; width: 150px; height: 150px; border-radius: 50%; margin: 6px auto 10px;
  background: conic-gradient(var(--c) calc(var(--p) * 1%), #e5e7eb 0); display:flex;
  align-items:center; justify-content:center;}
.ring .inner {width: 116px; height: 116px; border-radius: 50%; background: #fff; display:flex;
  flex-direction:column; align-items:center; justify-content:center;}
.ring .num {font-size: 2.3rem; font-weight: 800; color: #1F2937; line-height: 1;}
.ring .lab {font-size: .75rem; color:#6b7280; margin-top: 4px;}
.chip {display:inline-block; padding: 5px 12px; border-radius: 999px; font-size: .85rem;
  margin: 3px 4px 3px 0; font-weight: 500;}
.chip.good {background:#dcfce7; color:#166534;}
.chip.gap {background:#fee2e2; color:#991b1b;}
.chip.key {background:#e0e7ff; color:#3730a3;}
.card {background:#fff; border:1px solid #e5e7eb; border-radius:16px; padding: 18px 20px; margin-bottom: 14px;}
.card h4 {margin: 0 0 10px 0; color:#1F3A5F;}
.card ul {margin: 0; padding-left: 18px;}
.card li {margin-bottom: 6px; color:#374151;}
.paper {background:#fff; border-radius: 6px; padding: 34px 40px; margin-top: 12px;
  box-shadow: 0 8px 30px rgba(15,23,42,.12); border: 1px solid #e5e7eb; color:#111827;
  font-family: Calibri, "Segoe UI", Helvetica, Arial, sans-serif; font-size: .93rem; line-height: 1.45;}
.paper .nm {text-align:center; font-size: 1.7rem; font-weight: 800; color:#1F3A5F; letter-spacing:.5px;}
.paper .hl {text-align:center; font-weight: 700; margin-top: 2px;}
.paper .ct {text-align:center; color:#555; font-size:.82rem; margin-bottom: 4px;}
.paper .sh {color:#1F3A5F; font-weight:700; font-size:.88rem; letter-spacing:.8px; text-transform:uppercase;
  border-bottom: 1.5px solid #1F3A5F; margin: 14px 0 6px; padding-bottom: 1px;}
.paper .row {display:flex; justify-content:space-between; gap: 12px; margin-top: 6px;}
.paper .row .r {color:#555; white-space: nowrap;}
.paper .sub {color:#555; font-style: italic;}
.paper ul {margin: 2px 0 0 0; padding-left: 18px;}
.paper li {margin-bottom: 2px;}
.hint {color:#6b7280; font-size:.85rem;}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

st.session_state.setdefault("rv", 0)  # resume-uploader version
st.session_state.setdefault("jv", 0)  # job-fields version


def _new_job():
    """Keep the uploaded resume, clear the job fields and the result."""
    st.session_state.pop("result", None)
    st.session_state.pop("job_title", None)
    st.session_state["jv"] += 1


def _start_over():
    """Clear everything except the API key."""
    _new_job()
    st.session_state["rv"] += 1


st.markdown(
    """<div class="hero">
    <h1>📄 AI Resume Tailor</h1>
    <p>Paste a job description and get a resume rewritten for that exact role, ready as Word or PDF.</p>
    <span class="pill">✨ Tailored wording</span><span class="pill">🎯 ATS keywords</span>
    <span class="pill">📊 Skill gap analysis</span><span class="pill">⬇️ Word + PDF</span>
    </div>""",
    unsafe_allow_html=True,
)

with st.sidebar:
    st.markdown("### 🔑 Gemini API key")
    api_key = st.text_input(
        "API key", type="password", value=os.getenv("GEMINI_API_KEY", ""), label_visibility="collapsed"
    )
    st.caption("Get a free key at aistudio.google.com. Tip: set the GEMINI_API_KEY environment variable to skip typing it.")
    st.markdown("---")
    st.markdown("### 💡 Tips")
    st.markdown(
        "- Paste the **full** job description\n"
        "- Use a text-based PDF (not a scan)\n"
        "- Always read the result before sending\n"
        "- Run once per job you apply to"
    )

col_a, col_b = st.columns(2, gap="large")
with col_a:
    with st.container(border=True):
        st.markdown('<div class="step-title"><span class="step-badge">1</span>Your resume</div>', unsafe_allow_html=True)
        st.markdown('<div class="hint">Upload your current resume as a PDF.</div>', unsafe_allow_html=True)
        resume_file = st.file_uploader("Resume PDF", type=["pdf"], label_visibility="collapsed",
                                       key=f"resume_{st.session_state['rv']}")
with col_b:
    with st.container(border=True):
        st.markdown('<div class="step-title"><span class="step-badge">2</span>The job</div>', unsafe_allow_html=True)
        job_title = st.text_input("Target job title", placeholder="e.g. AI / Machine Learning Intern",
                                  key=f"title_{st.session_state['jv']}")
        job_desc = st.text_area("Job description", height=130, placeholder="Paste the full job description here...",
                                key=f"desc_{st.session_state['jv']}")

go = st.button("✨ Tailor my resume", type="primary", use_container_width=True)

if go:
    if not api_key:
        st.error("Enter your Gemini API key in the sidebar.")
    elif not resume_file or not job_title.strip() or not job_desc.strip():
        st.error("Please upload a resume and fill in the job title and description.")
    else:
        try:
            resume_text = extract_pdf_text(resume_file)
        except Exception as e:
            st.error(f"Couldn't read that PDF: {e}")
            st.stop()
        if len(resume_text) < 50:
            st.error("No text found in the PDF (it may be a scanned image). Try a text-based PDF.")
            st.stop()

        prompt = build_prompt(resume_text, job_title.strip(), job_desc.strip())
        try:
            with st.status("Tailoring your resume...", expanded=False) as status:
                result = analyze(api_key, prompt, status)
                status.update(label="Done!", state="complete")
            st.session_state["result"] = result
            st.session_state["job_title"] = job_title.strip()
        except errors.APIError as e:
            code = getattr(e, "code", None)
            if code in (401, 403):
                st.error("API key rejected. Check the key and its permissions.")
            else:
                st.error(f"Gemini API error ({code}): {e}")
            st.stop()
        except Exception as e:
            st.error(f"Gemini is busy or unreachable right now. Try again in a minute.\n\nDetails: {e}")
            st.stop()


def _e(s) -> str:
    return _html.escape(_clean(s or ""))


def _chips(items, kind) -> str:
    return "".join(f'<span class="chip {kind}">{_e(i)}</span>' for i in items)


def _bullets(items) -> str:
    return "<ul>" + "".join(f"<li>{_e(b)}</li>" for b in items) + "</ul>"


def resume_to_html(r: Resume) -> str:
    h = [f'<div class="nm">{_e(r.name).upper()}</div>', f'<div class="hl">{_e(r.headline)}</div>',
         f'<div class="ct">{" &nbsp;|&nbsp; ".join(_e(c) for c in r.contact if c.strip())}</div>',
         '<div class="sh">Professional Summary</div>', f"<div>{_e(r.summary)}</div>",
         '<div class="sh">Education</div>']
    for e in r.education:
        h.append(f'<div class="row"><div><b>{_e(e.degree)}</b>, <span class="sub">{_e(e.institution)}</span></div>'
                 f'<div class="r">{_e(e.dates)}</div></div>')
        if e.details.strip():
            h.append(f"<div>{_e(e.details)}</div>")
    h.append('<div class="sh">Technical Skills</div>')
    for g in r.skill_groups:
        h.append(f"<div><b>{_e(g.category)}:</b> {_e(', '.join(g.items))}</div>")
    h.append('<div class="sh">Projects</div>')
    for p in r.projects:
        tech = f' <span class="sub">| {_e(p.tech)}</span>' if p.tech.strip() else ""
        h.append(f'<div style="margin-top:6px"><b>{_e(p.name)}</b>{tech}</div>{_bullets(p.bullets)}')
    if r.experience:
        h.append('<div class="sh">Experience</div>')
        for x in r.experience:
            h.append(f'<div class="row"><div><b>{_e(x.title)}</b>, <span class="sub">{_e(x.organization)}</span></div>'
                     f'<div class="r">{_e(x.dates)}</div></div>{_bullets(x.bullets)}')
    if r.achievements:
        h.append('<div class="sh">Achievements</div>' + _bullets(r.achievements))
    if r.languages.strip():
        h.append(f'<div class="sh">Languages</div><div>{_e(r.languages)}</div>')
    return f'<div class="paper">{"".join(h)}</div>'


# Results live in session_state so downloads don't wipe the page
if "result" in st.session_state:
    result: Analysis = st.session_state["result"]
    r = result.resume
    base_name = f"{_slug(r.name)}_{_slug(st.session_state['job_title'])}"
    score = min(max(int(result.match_score), 0), 100)
    color = "#16a34a" if score >= 75 else "#f59e0b" if score >= 50 else "#dc2626"

    st.markdown("---")
    n1, n2, _sp = st.columns([1.4, 1, 2.6])
    n1.button("➕ Another job, same resume", on_click=_new_job, use_container_width=True, key="btn_newjob")
    n2.button("🔄 Start over", on_click=_start_over, use_container_width=True, key="btn_reset")
    tab_resume, tab_analysis = st.tabs(["📝 Tailored resume", "📊 Match analysis"])

    with tab_resume:
        try:
            docx_bytes, pdf_bytes = build_docx(r), build_pdf(r)
        except Exception as e:
            st.error(f"Couldn't build the files: {e}")
            docx_bytes = pdf_bytes = None

        top_l, top_r = st.columns([3, 2])
        with top_l:
            st.markdown(f"**Tailored for:** {_e(st.session_state['job_title'])}")
            st.markdown('<div class="hint">Read every line and make sure it is all true before you send it.</div>',
                        unsafe_allow_html=True)
        if docx_bytes and pdf_bytes:
            with top_r:
                d1, d2 = st.columns(2)
                d1.download_button(
                    "⬇️ Word", docx_bytes, f"{base_name}.docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True,
                )
                d2.download_button("⬇️ PDF", pdf_bytes, f"{base_name}.pdf", "application/pdf",
                                   use_container_width=True)
        st.markdown(resume_to_html(r), unsafe_allow_html=True)
        if result.ats_keywords:
            st.markdown("##### 🎯 Job keywords worked into your resume")
            st.markdown(_chips(result.ats_keywords, "key"), unsafe_allow_html=True)

    with tab_analysis:
        c_ring, c_sum = st.columns([1, 2], gap="large")
        with c_ring:
            st.markdown(
                f'<div class="ring" style="--p:{score};--c:{color}"><div class="inner">'
                f'<div class="num">{score}</div><div class="lab">MATCH SCORE</div></div></div>',
                unsafe_allow_html=True,
            )
        with c_sum:
            st.markdown(f'<div class="card"><h4>Summary</h4>{_e(result.summary)}</div>', unsafe_allow_html=True)

        c1, c2 = st.columns(2, gap="large")
        with c1:
            st.markdown(f'<div class="card"><h4>✅ Strengths</h4>{_chips(result.strengths, "good")}</div>',
                        unsafe_allow_html=True)
        with c2:
            st.markdown(f'<div class="card"><h4>⚠️ Missing skills</h4>{_chips(result.missing_skills, "gap")}</div>',
                        unsafe_allow_html=True)
        st.markdown(f'<div class="card"><h4>💡 How to improve your chances</h4>{_bullets(result.recommendations)}</div>',
                    unsafe_allow_html=True)